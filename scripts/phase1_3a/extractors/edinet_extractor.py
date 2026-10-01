"""
scripts/phase1_3a/extractors/edinet_extractor.py
Phase 1.3A Extractor for FSA EDINET Taxonomy with Section 14 Composite Taxonomy Label Detection.
Extracts accounting line items, statements, and financial taxonomy concepts across industry sheets.
"""

from pathlib import Path
from typing import List, Tuple
import re
import openpyxl
from datetime import datetime, timezone

from scripts.phase1_3a.extractors.base import BasePhase13Extractor
from scripts.phase1_3a.models import RawCandidate


def clean_label(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r"\[.*?\]", "", text).strip()
    return text


def classify_edinet_label(label: str) -> Tuple[str, bool]:
    """
    Classifies an EDINET label per Phase 1.3A Section 14:
    Returns (classification, is_composite_reporting_label).
    Classifications:
    - CANONICAL_CONCEPT
    - COMPOSITE_REPORTING_LABEL
    - TAXONOMY_VARIANT
    - POSSIBLE_CANONICAL
    - NEEDS_REVIEW
    """
    if not label:
        return ("NEEDS_REVIEW", True)

    # Pure navigation / structural artifacts
    if label in ("目次", "勘定科目リストについて", "用語一覧", "要素名", "標準ラベル"):
        return ("TAXONOMY_VARIANT", True)

    # Explicitly known atomic professional concepts (preserve these!)
    atomic_known = {
        "契約資産", "受取手形", "売掛金", "買掛金", "支払手形", "コールローン", "買入手形",
        "売戻条件付買現先勘定", "債権代位弁済", "当座預金", "定期預金", "前払金", "前受金",
        "前払費用", "未払費用", "未収収益", "未払金", "未収入金", "短期借入金", "長期借入金",
        "減価償却累計額", "貸倒引当金", "賞与引当金", "役員賞与引当金", "退職給付引当金",
        "資本金", "資本準備金", "利益準備金", "繰越利益剰余金", "自己株式"
    }
    if label in atomic_known:
        return ("CANONICAL_CONCEPT", False)

    # Net valuation indicators
    has_net = bool(re.search(r"[\(（]純額[\)）]", label))

    # Composite coordinator keywords
    has_and = "及び" in label or "、" in label or "又は" in label
    has_other = label.startswith("その他の") or label.startswith("その他")
    has_etc = label.endswith("等")

    # Composite reporting label detection (Section 14)
    # e.g. "受取手形、売掛金及び契約資産", "受取手形及び売掛金", "受取手形及び売掛金(純額)",
    # "売掛金及び契約資産", "売掛金及び契約資産(純額)", "受取手形(純額)", "売掛金(純額)", "契約資産(純額)", "コールローン及び買入手形"
    if (has_and and has_net) or (has_and and ("及び" in label or "、" in label)):
        return ("COMPOSITE_REPORTING_LABEL", True)
    if has_net:
        return ("COMPOSITE_REPORTING_LABEL", True)
    if has_other and len(label) >= 6:
        return ("COMPOSITE_REPORTING_LABEL", True)
    if has_and:
        return ("COMPOSITE_REPORTING_LABEL", True)

    if has_etc:
        return ("TAXONOMY_VARIANT", False)

    return ("CANONICAL_CONCEPT", False)


class FsaEdinetPhase13Extractor(BasePhase13Extractor):
    def extract_candidates(self) -> List[RawCandidate]:
        candidates = []
        seen = set()
        now_iso = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

        # 1. 1f_AccountList.xlsx (Financial statement line items across all industries)
        account_list_path = self.raw_snapshot_path / "1f_AccountList.xlsx"
        if account_list_path.exists():
            wb = openpyxl.load_workbook(account_list_path, data_only=True, read_only=True)
            for sheet_name in wb.sheetnames:
                if sheet_name in ("目次", "勘定科目リストについて"):
                    continue
                ws = wb[sheet_name]

                # Determine domain / subdomain mapping per industry
                if sheet_name in ("銀行・信託業", "銀行・信託業（特定取引勘定設置銀行）", "第一種金融商品取引業", "損害保険業", "生命保険業", "投資運用業", "投資業", "特定金融業", "商品先物取引業", "資産流動化業", "投資信託受益証券"):
                    primary_dom = "finance"
                    subdom = "banking" if "銀行" in sheet_name else "capital_markets"
                elif sheet_name == "海運事業":
                    primary_dom = "trade"
                    subdom = "shipping"
                elif sheet_name in ("電気通信事業", "鉄道事業", "電気事業", "ガス事業", "高速道路事業", "リース事業"):
                    primary_dom = "business"
                    subdom = "industry_operations"
                else:
                    primary_dom = "accounting"
                    subdom = "financial_accounting"

                for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    if r_idx <= 2 or len(row) < 9:
                        continue

                    ja_label = str(row[1]).strip() if row[1] is not None else ""
                    en_label = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
                    element_name = str(row[8]).strip() if len(row) > 8 and row[8] is not None else ""
                    is_abstract = str(row[13]).strip().lower() if len(row) > 13 and row[13] is not None else "false"

                    clean_ja = clean_label(ja_label)
                    is_statement = clean_ja in [
                        "貸借対照表", "損益計算書", "株主資本等変動計算書", "包括利益計算書", "キャッシュ・フロー計算書"
                    ]

                    if is_abstract == "true" and not is_statement:
                        continue
                    if not clean_ja or clean_ja in seen:
                        continue
                    if len(clean_ja) < 2 or len(clean_ja) > 50:
                        continue

                    seen.add(clean_ja)
                    seq = len(candidates) + 1
                    cand_id = f"cand-edinet-{seq:06d}"
                    record_id = element_name if element_name else f"edinet_acc_{r_idx}"
                    locator = f"sheet:{sheet_name}, row:{r_idx}, element:{element_name}"

                    classification, is_composite = classify_edinet_label(clean_ja)
                    context = f"EDINET Taxonomy account item. Official EN: {en_label}" if en_label else "EDINET Taxonomy account item."
                    if is_composite:
                        context += f" [XBRL_CLASSIFICATION:{classification}]"

                    candidates.append(RawCandidate(
                        candidate_id=cand_id,
                        source_id=self.source_id,
                        source_version=self.source_version,
                        source_record_id=record_id,
                        source_locator=locator,
                        source_term_exact=ja_label,
                        source_context=context,
                        source_definition=None,
                        extracted_at=now_iso,
                        extractor_version=self.extractor_version,
                        raw_snapshot_hash=self.raw_snapshot_hash,
                        primary_domain=primary_dom,
                        subdomain=subdom,
                        authority_class=self.authority_class,
                        reuse_status=self.reuse_status
                    ))
            wb.close()

        # 2. 1e_ElementList.xlsx (Corporate governance, statutory filing, and audit elements)
        element_list_path = self.raw_snapshot_path / "1e_ElementList.xlsx"
        if element_list_path.exists():
            wb_el = openpyxl.load_workbook(element_list_path, data_only=True, read_only=True)
            for sheet_name in wb_el.sheetnames:
                if sheet_name in ("目次", "タクソノミ要素リストについて"):
                    continue
                ws = wb_el[sheet_name]
                for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    if r_idx <= 2 or len(row) < 4:
                        continue
                    ja_label = str(row[0]).strip() if row[0] is not None else ""
                    en_label = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
                    element_name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""

                    clean_ja = clean_label(ja_label)
                    if not clean_ja or clean_ja in seen:
                        continue
                    if len(clean_ja) < 2 or len(clean_ja) > 50:
                        continue

                    # Filter out purely generic English/code rows
                    if re.match(r"^[A-Za-z0-9_\-\.\s]+$", clean_ja):
                        continue

                    seen.add(clean_ja)
                    seq = len(candidates) + 1
                    cand_id = f"cand-edinet-{seq:06d}"
                    record_id = element_name if element_name else f"edinet_el_{r_idx}"
                    locator = f"file:1e_ElementList.xlsx, sheet:{sheet_name}, row:{r_idx}"

                    classification, is_composite = classify_edinet_label(clean_ja)
                    context = f"EDINET Taxonomy disclosure element. Official EN: {en_label}" if en_label else "EDINET Taxonomy disclosure element."
                    if is_composite:
                        context += f" [XBRL_CLASSIFICATION:{classification}]"

                    # Assign legal/corporate or accounting domain
                    if any(kw in clean_ja for kw in ("ガバナンス", "役員", "株主", "総会", "定款", "取締役", "監査役", "提出者")):
                        p_dom = "legal"
                        s_dom = "corporate_governance"
                    elif any(kw in clean_ja for kw in ("監査", "独立監査人", "監査意見")):
                        p_dom = "accounting"
                        s_dom = "audit"
                    else:
                        p_dom = "accounting"
                        s_dom = "financial_reporting"

                    candidates.append(RawCandidate(
                        candidate_id=cand_id,
                        source_id=self.source_id,
                        source_version=self.source_version,
                        source_record_id=record_id,
                        source_locator=locator,
                        source_term_exact=ja_label,
                        source_context=context,
                        source_definition=None,
                        extracted_at=now_iso,
                        extractor_version=self.extractor_version,
                        raw_snapshot_hash=self.raw_snapshot_hash,
                        primary_domain=p_dom,
                        subdomain=s_dom,
                        authority_class=self.authority_class,
                        reuse_status=self.reuse_status
                    ))
            wb_el.close()

        return candidates
