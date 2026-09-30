#!/usr/bin/env python3
"""
scripts/extract_fsa.py
Extracts official accounting terms from FSA EDINET 2026 Taxonomy (1f_AccountList.xlsx).
Filters out abstract section headers, extracts standard labels, English mappings,
balance types, and accounting line items.
Outputs: data/extracted/fsa_accounting_candidates.jsonl
"""

import json
import re
from pathlib import Path
import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
INPUT_EXCEL = BASE_DIR / "data" / "raw" / "fsa" / "edinet" / "2026" / "1f_AccountList.xlsx"
OUTPUT_FILE = BASE_DIR / "data" / "extracted" / "fsa_accounting_candidates.jsonl"


def clean_label(text: str) -> str:
    if not text:
        return ""
    # Remove EDINET bracket annotations like [目次項目], [タイトル項目], [軸]
    text = re.sub(r"\[.*?\]", "", text).strip()
    return text


def extract_accounting_terms():
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    print(f"[*] Extracting FSA accounting terms from {INPUT_EXCEL.name} ...")

    wb = openpyxl.load_workbook(INPUT_EXCEL, data_only=True)
    sheets_to_process = ["一般商工業", "建設業", "銀行・信託業"]

    extracted = []
    seen_surfaces = set()

    for sheet_name in sheets_to_process:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        print(f"    Processing sheet: {sheet_name} ({ws.max_row} rows) ...")

        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if r_idx <= 2:
                continue  # header rows

            # Columns in 1f_AccountList.xlsx:
            # 0: 科目分類
            # 1: 標準ラベル（日本語）
            # 2: 冗長ラベル（日本語）
            # 3: 標準ラベル（英語）
            # 4: 冗長ラベル（英語）
            # 5: 用途区分
            # 8: 要素名
            # 11: periodType
            # 12: balance
            # 13: abstract (true/false)

            ja_label = str(row[1]).strip() if row[1] is not None else ""
            en_label = str(row[3]).strip() if row[3] is not None else ""
            element_name = str(row[8]).strip() if len(row) > 8 and row[8] is not None else ""
            balance = str(row[12]).strip() if len(row) > 12 and row[12] is not None else ""
            is_abstract = str(row[13]).strip().lower() if len(row) > 13 and row[13] is not None else "false"

            clean_ja = clean_label(ja_label)

            # Include financial statements as top-level accounts
            is_statement = clean_ja in ["貸借対照表", "損益計算書", "株主資本等変動計算書", "包括利益計算書", "キャッシュ・フロー計算書"]
            if is_abstract == "true" and not is_statement:
                continue
            if not clean_ja or clean_ja in seen_surfaces:
                continue

            seen_surfaces.add(clean_ja)
            cand_seq = len(extracted) + 1
            extracted.append({
                "extracted_candidate_id": f"cand-fsa-{cand_seq:06d}",
                "source_id": "fsa_edinet_2026",
                "source_file": "1f_AccountList.xlsx",
                "source_record_id": element_name if element_name else f"fsa_row_{r_idx}",
                "source_term_exact": ja_label,
                "surface_candidate": clean_ja,
                "official_en": en_label,
                "element_name": element_name,
                "balance": balance,
                "sheet": sheet_name,
                "category": "accounting",
                "is_statement_title": is_statement
            })

    wb.close()

    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        for item in extracted:
            f.write(json.dumps(item, ensure_ascii=False) + "\n")

    print(f"[+] Successfully extracted {len(extracted)} FSA accounting candidates to {OUTPUT_FILE.relative_to(BASE_DIR)}")


if __name__ == "__main__":
    extract_accounting_terms()
