"""
scripts/source_engine/extractors/edinet.py
Extractor for FSA EDINET Taxonomy (1f_AccountList.xlsx).
Extracts accounting line items, statements, and financial taxonomy concepts.
"""

from pathlib import Path
from typing import List
import re
import openpyxl
from datetime import datetime, timezone

from scripts.source_engine.extractors.base import BaseExtractor
from scripts.source_engine.models import ExtractedCandidate


def clean_label(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r"\[.*?\]", "", text).strip()
    return text


class FsaEdinetExtractor(BaseExtractor):
    def extract_candidates(self) -> List[ExtractedCandidate]:
        excel_path = self.raw_snapshot_path / "1f_AccountList.xlsx"
        if not excel_path.exists():
            return []

        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheets = ["一般商工業", "建設業", "銀行・信託業"]
        candidates = []
        seen = set()

        now_iso = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

        for sheet_name in sheets:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]

            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if r_idx <= 2:
                    continue

                ja_label = str(row[1]).strip() if row[1] is not None else ""
                en_label = str(row[3]).strip() if row[3] is not None else ""
                element_name = str(row[8]).strip() if len(row) > 8 and row[8] is not None else ""
                is_abstract = str(row[13]).strip().lower() if len(row) > 13 and row[13] is not None else "false"

                clean_ja = clean_label(ja_label)
                is_statement = clean_ja in [
                    "貸借対照表", "損益計算書", "株主資本等変動計算書", "包括利益計算書", "キャッシュ・フロー計算書"
                ]

                if is_abstract == "true" and not is_statement:
                    continue
                if not clean_ja or clean_ja in seen:
                    continue
                if len(clean_ja) < 2 or len(clean_ja) > 40:
                    continue

                seen.add(clean_ja)
                seq = len(candidates) + 1
                cand_id = f"cand-fsa-{seq:06d}"
                record_id = element_name if element_name else f"fsa_row_{r_idx}"
                locator = f"sheet:{sheet_name}, row:{r_idx}, element:{element_name}"
                context = f"EDINET Taxonomy account item. Official EN: {en_label}" if en_label else "EDINET Taxonomy account item."

                primary_dom = "finance" if sheet_name == "銀行・信託業" else "accounting"
                if primary_dom == "finance":
                    subdom = "banking"
                elif is_statement:
                    subdom = "financial_statements"
                else:
                    subdom = "financial_accounting"

                candidates.append(ExtractedCandidate(
                    candidate_id=cand_id,
                    source_id=self.source_id,
                    source_version=self.source_version,
                    source_record_id=record_id,
                    source_locator=locator,
                    source_term_exact=ja_label,
                    source_context=context,
                    source_definition=None,
                    extracted_at=now_iso,
                    extractor_version=self.extractor_version,
                    raw_snapshot_hash=self.raw_snapshot_hash,
                    primary_domain=primary_dom,
                    subdomain=subdom,
                    authority_class=self.authority_class,
                    reuse_status=self.reuse_status
                ))

        wb.close()
        return candidates
